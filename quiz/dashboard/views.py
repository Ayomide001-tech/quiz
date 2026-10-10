from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from account.models import Account
import json
import logging
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods, require_POST
from .models import Question, category, QuizAttempt
from openpyxl import Workbook, load_workbook
from django.urls import reverse_lazy
from django.views.generic import UpdateView, DeleteView
from django.conf import settings
from django.db import transaction
from google import genai
from google.genai import types

logger = logging.getLogger(__name__)




def upload_questions(request):
    return render(request, 'questioncenter/upload_questions.html')

def download_question_template(request):
    # Create a new workbook and select the active worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    headers =  [
        'category',
        "question",
        "option_a",
        "option_b",
        "option_c",
        "option_d",
        "correct_answer",

    ]

    ws.append(headers)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=question_template.xlsx'
    wb.save(response)
    return response







# --------------------------------------------------------- How to import questions from excel file--------------------------------------------------------


def import_questions_from_excel(request):
    if request.method != "POST":
        return redirect("dashboard")

    uploaded_file = request.FILES.get("excel_file")
    if not uploaded_file:
        messages.error(request, "Please choose an Excel file before clicking Send.")
        return redirect("dashboard")

    if not uploaded_file.name.lower().endswith(".xlsx"):
        messages.error(request, "This file type is not supported. Please upload an .xlsx Excel file.")
        return redirect("dashboard")

    try:
        workbook = load_workbook(uploaded_file, data_only=True)
        sheet = workbook.active

        required_headers = [
            "category",
            "question",
            "option_a",
            "option_b",
            "option_c",
            "option_d",
            "correct_answer",
        ]

        first_row = [cell.value for cell in sheet[1]]
        if first_row != required_headers:
            messages.error(
                request,
                "We could not read this file. Please download a new template and keep the first row unchanged."
            )
            return redirect("dashboard")

        created = 0
        skipped = 0
        errors = []

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if not row or all(cell is None for cell in row):
                continue

            category_name, question, option_a, option_b, option_c, option_d, correct_answer = row[:7]

            if not all([category_name, question, option_a, option_b, option_c, option_d, correct_answer]):
                skipped += 1
                errors.append(f"Question {row_number - 1} was not added because one or more cells are empty.")
                continue

            category_name = " ".join(str(category_name).split())
            question = " ".join(str(question).split())
            correct_answer = str(correct_answer).strip().upper()

            if correct_answer not in {"A", "B", "C", "D"}:
                skipped += 1
                errors.append(
                    f"Question {row_number - 1} was not added because the correct answer must be A, B, C, or D. Lowercase letters such as b are accepted."
                )
                continue

            category_obj = category.objects.filter(
                category_name__iexact=category_name
            ).first()
            if category_obj is None:
                category_obj = category.objects.create(category_name=category_name)

            duplicate_exists = Question.objects.filter(
                category=category_obj,
                question__iexact=question,
            ).exists()
            if duplicate_exists:
                skipped += 1
                errors.append(f"Question {row_number - 1} was not added because it already exists in this category.")
                continue

            Question.objects.create(
                category=category_obj,
                question=question,
                option_a=str(option_a).strip(),
                option_b=str(option_b).strip(),
                option_c=str(option_c).strip(),
                option_d=str(option_d).strip(),
                correct_answer=correct_answer,
            )
            created += 1

        if created:
            messages.success(request, f"Done! {created} question(s) were added successfully.")
        if skipped:
            messages.warning(request, f"Upload finished, but {skipped} question(s) were not added.")
            for error in errors[:5]:
                messages.error(request, error)
            if len(errors) > 5:
                messages.error(request, f"And {len(errors) - 5} more row error(s).")
        if not created and not skipped:
            messages.warning(request, "The file was uploaded, but it contains no questions to add.")
        return redirect("dashboard")

    except Exception as e:
        messages.error(request, "We could not process this file. Please check that it is a valid Excel template and try again.")
        return redirect("dashboard")





# ------------------------------------------------------------End of How to import questions from excel file--------------------------------------------------------






# -------------------------------------------------AI Question--------------------------------------------------------


def _generate_ai_questions(topic, difficulty, question_count):
    api_key = settings.GEMINI_API_KEY
    if not api_key:
        raise ValueError('The AI service is not configured yet.')

    prompt = f'''Create exactly {question_count} {difficulty} multiple-choice quiz question(s) about {topic}.
Return only a JSON array. Each item must contain exactly these keys:
question, option_a, option_b, option_c, option_d, correct_answer.
correct_answer must be one uppercase letter: A, B, C, or D.
Do not repeat questions. Do not include markdown or explanations.'''

    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model='gemini-3.6-flash',
        contents=prompt,
        config=types.GenerateContentConfig(response_mime_type='application/json'),
    )
    if not response.text:
        raise ValueError('Gemini returned an empty response.')

    generated_questions = json.loads(response.text)
    if not isinstance(generated_questions, list):
        raise ValueError('The AI returned an invalid question list.')

    valid_questions = []
    for generated in generated_questions:
        required_fields = ['question', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_answer']
        if not isinstance(generated, dict) or any(
            not str(generated.get(field, '')).strip() for field in required_fields
        ):
            continue

        question_text = ' '.join(str(generated['question']).split())
        correct_answer = str(generated['correct_answer']).strip().upper()
        if correct_answer not in {'A', 'B', 'C', 'D'}:
            continue
        if any(item['question'].casefold() == question_text.casefold() for item in valid_questions):
            continue

        valid_questions.append({
            'question': question_text,
            'option_a': str(generated['option_a']).strip(),
            'option_b': str(generated['option_b']).strip(),
            'option_c': str(generated['option_c']).strip(),
            'option_d': str(generated['option_d']).strip(),
            'correct_answer': correct_answer,
        })

    return generated_questions, valid_questions


def generate_questions_with_ai(request):
    if not request.user.is_authenticated or not request.user.is_superuser:
        return redirect('homepage')

    categories = category.objects.all().order_by('category_name')
    if request.method == 'GET':
        return render(request, 'questioncenter/ai_generated.html', {'categories': categories})

    category_id = request.POST.get('category')
    difficulty = request.POST.get('difficulty', 'easy').strip().lower()
    try:
        question_count = int(request.POST.get('question_count', '1').strip())
    except ValueError:
        messages.error(request, 'Please enter a whole number of questions.')
        return redirect('generate_questions')

    selected_category = get_object_or_404(category, id=category_id)
    if difficulty not in {'easy', 'medium', 'hard'} or not 1 <= question_count <= 20:
        messages.error(request, 'Choose a valid difficulty and between 1 and 20 questions.')
        return redirect('generate_questions')

    try:
        generated_questions, valid_questions = _generate_ai_questions(
            selected_category.category_name, difficulty, question_count
        )
        created = 0
        for item in valid_questions:
            if Question.objects.filter(category=selected_category, question__iexact=item['question']).exists():
                continue
            Question.objects.create(category=selected_category, **item)
            created += 1
        messages.success(request, f'{created} question(s) were added to {selected_category.category_name}.')
        skipped = len(generated_questions) - created
        if skipped:
            messages.warning(request, f'{skipped} question(s) were skipped because they were duplicates or invalid.')
    except (json.JSONDecodeError, ValueError) as error:
        messages.error(request, str(error))
    except Exception as error:
        logger.exception('Gemini question generation failed: %s', error)
        messages.error(request, 'Gemini could not generate the questions. Please try again later.')
    return redirect('generate_questions')


def category_question_with_ai(request):
    if not request.user.is_authenticated or not request.user.is_superuser:
        return redirect('homepage')

    if request.method == 'GET':
        return render(request, 'questioncenter/category_with_ai.html')

    category_name = ' '.join(request.POST.get('category_name', '').split())
    try:
        quiz_time = int(request.POST.get('time', '60').strip())
        question_count = int(request.POST.get('question_count', '1').strip())
    except ValueError:
        messages.error(request, 'Quiz time and question count must be whole numbers.')
        return redirect('category_with_ai')

    difficulty = request.POST.get('difficulty', 'easy').strip().lower()
    if not category_name or not 1 <= len(category_name) <= 30:
        messages.error(request, 'Enter a category name between 1 and 30 characters.')
        return redirect('category_with_ai')
    if category.objects.filter(category_name__iexact=category_name).exists():
        messages.error(request, 'That category already exists. Use the existing-category AI page instead.')
        return redirect('category_with_ai')
    if not 1 <= quiz_time <= 600 or not 1 <= question_count <= 20 or difficulty not in {'easy', 'medium', 'hard'}:
        messages.error(request, 'Choose valid quiz time, question count, and difficulty values.')
        return redirect('category_with_ai')

    try:
        generated_questions, valid_questions = _generate_ai_questions(category_name, difficulty, question_count)
        if not valid_questions:
            messages.error(request, 'Gemini returned no valid questions, so the category was not created.')
            return redirect('category_with_ai')

        with transaction.atomic():
            new_category = category.objects.create(category_name=category_name, time=quiz_time)
            for item in valid_questions:
                Question.objects.create(category=new_category, **item)

        messages.success(request, f'Category {category_name} and {len(valid_questions)} question(s) were created.')
        if len(generated_questions) > len(valid_questions):
            messages.warning(request, 'Some AI questions were skipped because they were invalid or repeated.')
    except (json.JSONDecodeError, ValueError) as error:
        messages.error(request, str(error))
    except Exception as error:
        logger.exception('Gemini category creation failed: %s', error)
        messages.error(request, 'Gemini could not create the category questions. Please try again later.')
    return redirect('category_with_ai')

# ------------------------------------------------------------End of AI Question--------------------------------------------------------





def category_questions(request, category_id):
    cat = category.objects.get(id=category_id)
    questions = Question.objects.filter(category=cat)
    return render(request, 'dashboard/questions.html', {'questions': questions, 'category': cat})

# def display_question(request, category_id):
    
def display_questions(request, categorys_id):
    context = {}
    cat = category.objects.get(id=categorys_id)
    questions = Question.objects.filter(category=cat)
    context['questions'] = questions
    context['category'] = cat
    return render(request, 'questioncenter/viewquestion.html', context)





def viewaddedquestion(request):
    context = {}
    Questions = Question.objects.all()
    categorys = category.objects.all()
    context['categorys'] = categorys
    context['Questions'] = Questions


    return render(request, 'questioncenter/viewaddedquestion.html', context)

    

def dashboard(request):
    context = {}
    Questions = Question.objects.all()
    categorys = category.objects.all()
    context['categorys'] = categorys
    context['questions'] = Questions
    return render(request, 'dashboard/dashboard.html', context)


def homepage(request):
    return render(request, 'dashboard/homepage.html')

def superuserpage(request):
    if not request.user.is_authenticated:
       return redirect ('homepage')
    return render(request, 'superuser&admin/superuserpage.html', {})

def adminpage(request):
    return render(request, 'superuser&admin/adminpage.html', {})


@csrf_exempt
@require_http_methods(["POST"])
def create_categories_api(request):
    if not request.user.is_authenticated:
        return JsonResponse(
            {"success": False, "error": "You must be logged in."},
            status=401,
        )

    if not request.user.is_superuser:
        return JsonResponse(
            {"success": False, "error": "You are not allowed to create categories."},
            status=403,
        )

    try:
        data = json.loads(request.body)
        category_name = data.get('category_name')
        time = data.get('time')

        if not category_name or not time:
            return JsonResponse(
                {"success": False, "error": "All fields are required."},
                status=400,
            )

        # Check if category already exists
        if category.objects.filter(category_name=category_name).exists():
            return JsonResponse(
                {"success": False, "error": "This category already exists."},
                status=400,
            )

        new_category = category.objects.create(
            category_name=category_name,
            time=time,
        )

        return JsonResponse({
            "success": True,
            "message": "Category created successfully.",
            "new_category": {
                "id": new_category.id,
                "category_name": new_category.category_name,
                "time": new_category.time,
            },
        }, status=201)
    except json.JSONDecodeError:
        return JsonResponse(
            {"success": False, "error": "Invalid JSON."},
            status=400
        )
    except Exception as e:
        return JsonResponse(
            {"success": False, "error": str(e)},
            status=500
        )


def create_categories_page(request):
    return render(request, 'all_categories/create_categories_page.html')


def create_question(request):
    if not request.user.is_authenticated:
        return redirect('homepage')

    if not request.user.is_admin or not  request.user.is_staff or not request.user.is_superuser:
        return redirect('dashboard')

    categories = category.objects.all()
    context = {'categories': categories}
    return render(request, 'questioncenter/create_question.html', context)

@csrf_exempt
@require_http_methods(["POST"])
def create_question_api(request):
    if not request.user.is_authenticated:
        return JsonResponse(
            {"success": False, "error": "You must be logged in."},
            status=401,
        )

    if not request.user.is_admin or not  request.user.is_staff or not request.user.is_superuser:
        return JsonResponse(
            {"success": False, "error": "You are not allowed to create questions."},
            status=403,
        )

    try:
        data = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse(
            {"success": False, "error": "Invalid JSON."},
            status=400
        )

    categories = category.objects.all()
    
    category_id = data.get('category')
    question = data.get('question')
    option_a = data.get('option_a')
    option_b = data.get('option_b')
    option_c = data.get('option_c')
    option_d = data.get('option_d')
    correct_answer = data.get('correct_answer')

    if not category_id or not question or not option_a or not option_b or not option_c or not option_d or not correct_answer:
        return JsonResponse(
            {"success": False, "error": "All fields are required."},
            status=400,
        )

    try:
        category_instance = category.objects.get(id=category_id)

        new_question = Question.objects.create(
            category=category_instance,
            question=question,
            option_a=option_a,
            option_b=option_b,
            option_c=option_c,
            option_d=option_d,
            correct_answer=correct_answer
        )

        return JsonResponse({
            "success": True,
            "message": "Question created successfully.",
            "new_question": {
                "id": new_question.id,
                "category": new_question.category.id,
                "question": new_question.question,
                "option_a": new_question.option_a,
                "option_b": new_question.option_b,
                "option_c": new_question.option_c,
                "option_d": new_question.option_d,
                "correct_answer": new_question.correct_answer
            }
        })

    except category.DoesNotExist:
        return JsonResponse(
            {"success": False, "error": "Invalid category ID."},
            status=400,
        )
    except Exception as e:
        return JsonResponse(
            {"success": False, "error": str(e)},
            status=500
        )

        



@login_required
@require_POST
def save_score(request):
    score = int(request.POST.get('score', 0))
    category_id = request.POST.get('category_id')
    total_question = int(request.POST.get('total_question', 0) or 0)

    category_obj = None
    if category_id:
        category_obj = category.objects.filter(id=category_id).first()

    attempt = QuizAttempt.objects.create(
        user=request.user,
        category=category_obj,
        score=score,
        total_question=total_question,
    )

    # Keep the last score on the user model as a quick reference, but don't overwrite history.
    request.user.score = score
    request.user.save(update_fields=['score'])

    return JsonResponse({'success': True, 'score': score, 'attempt_id': attempt.id})


def viewScores(request):
    context = {}

    attempts = QuizAttempt.objects.select_related('user', 'category').order_by('-completed')
    user_attempts = []
    seen_users = set()

    for attempt in attempts:
        user_id = attempt.user_id
        if user_id not in seen_users:
            seen_users.add(user_id)
            user_attempts.append({
                'user': attempt.user,
                'latest_score': attempt.score,
                'attempts': [attempt],
            })
        else:
            for item in user_attempts:
                if item['user'].id == user_id:
                    item['attempts'].append(attempt)
                    item['latest_score'] = attempt.score
                    break

    context['user_attempts'] = user_attempts
    context['user_details'] = Account.objects.all()

    if request.user.is_authenticated:
        context['current_user'] = request.user
        latest_attempt = QuizAttempt.objects.filter(user=request.user).order_by('-completed').first()
        context['current_user_score'] = latest_attempt.score if latest_attempt else request.user.score

    return render(request, "dashboard/viewScores.html", context)






# ---------------------------------------- update category and delete category -----------------------------------


class categoryUpdateView(UpdateView):
    model = category
    fields = ['category_name', 'time']
    template_name = 'all_categories/update_category.html'
    success_url = reverse_lazy('viewaddedquestion')

    def form_valid(self, form):
        messages.success(self.request, "Category updated successfully.")
        return super().form_valid(form)

class categoryDeleteView(DeleteView):
    model = category
    template_name = 'all_categories/delete_category.html'
    success_url = reverse_lazy('viewaddedquestion')

    def delete(self, request, *args, **kwargs):
        response = super().delete(request, *args, **kwargs)
        messages.success(self.request, "Category deleted successfully.")
        return response



# --------------------------------------update question and delete question--------------------------------

class QuestionUpdateView(UpdateView):
    model = Question
    fields = ['category', 'question', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_answer']
    template_name = 'all_categories/update_question.html'
    success_url = reverse_lazy('viewaddedquestion')

    def form_valid(self, form):
        old_category = self.get_object().category
        response = super().form_valid(form)
        messages.success(self.request, "Question updated successfully.")

        if old_category.questions.count() == 0:
            old_category.delete()

        return response

class QuestionDeleteView(DeleteView):
    model = Question
    template_name = 'questioncenter/delete_question.html'
    success_url = reverse_lazy('viewaddedquestion')

    def delete(self, request, *args, **kwargs):
        question = self.get_object()
        question_category = question.category
        response = super().delete(request, *args, **kwargs)

        if question_category.questions.count() == 0:
            question_category.delete()

        messages.success(self.request, "Question deleted successfully.")
        return response
